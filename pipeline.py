#!/usr/bin/env python3
"""Hitpoint Sales Intelligence Agent — 批量管线入口。

  输入:单个公司 / CSV(1000家) / Excel(100家)
  流程:读取公司列表 → 研究引擎调研 → 机会评分 → 排序 → 输出 Excel + 每家 Markdown 情报

用法:
  python pipeline.py --company Medtronic --engine manual
  python pipeline.py --input data/input_sample.csv --engine manual --out output
  python pipeline.py --input accounts.xlsx --engine api --workers 4

研究引擎:
  manual  读取 data/enriched/<slug>.json(人工 / Claude Code 内置搜索预调研)
  api     用 Claude API + web_search 自动调研(需 ANTHROPIC_API_KEY、pip install anthropic)
"""

from __future__ import annotations

import argparse
import csv
import json
import os
import sys
from concurrent.futures import ThreadPoolExecutor, as_completed

import report
from engines import get_engine
from scoring import score_record
from util import slugify

COMPANY_COLUMN_HINTS = ["company", "company name", "name", "公司", "公司名称", "客户", "account"]


# ---------- 读取输入 ----------

def read_companies(args) -> list[str]:
    if args.company:
        return [args.company]
    if not args.input:
        sys.exit("请用 --company 或 --input 指定输入。")
    path = args.input
    if not os.path.exists(path):
        sys.exit(f"找不到输入文件:{path}")
    ext = os.path.splitext(path)[1].lower()
    if ext in (".csv", ".tsv", ".txt"):
        return _read_delimited(path, "\t" if ext == ".tsv" else ",")
    if ext in (".xlsx", ".xlsm"):
        return _read_excel(path)
    sys.exit(f"不支持的输入格式:{ext}(用 .csv / .tsv / .xlsx)")


def _pick_column(header: list[str]) -> int:
    lower = [str(h or "").strip().lower() for h in header]
    for hint in COMPANY_COLUMN_HINTS:
        if hint in lower:
            return lower.index(hint)
    return 0  # 找不到就用第一列


def _read_delimited(path: str, delim: str) -> list[str]:
    with open(path, newline="", encoding="utf-8-sig") as f:
        rows = list(csv.reader(f, delimiter=delim))
    if not rows:
        return []
    col = _pick_column(rows[0])
    has_header = any(h.strip().lower() in COMPANY_COLUMN_HINTS for h in rows[0])
    data = rows[1:] if has_header else rows
    return _dedupe([r[col].strip() for r in data if len(r) > col and r[col].strip()])


def _read_excel(path: str) -> list[str]:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = [[c if c is not None else "" for c in row] for row in ws.iter_rows(values_only=True)]
    if not rows:
        return []
    header = [str(c) for c in rows[0]]
    col = _pick_column(header)
    has_header = any(str(h).strip().lower() in COMPANY_COLUMN_HINTS for h in header)
    data = rows[1:] if has_header else rows
    return _dedupe([str(r[col]).strip() for r in data if len(r) > col and str(r[col]).strip()])


def _dedupe(names: list[str]) -> list[str]:
    return list(dict.fromkeys(names))


# ---------- 主流程 ----------

def run(args) -> None:
    companies = read_companies(args)
    if not companies:
        sys.exit("输入里没有公司名。")
    print(f"📋 共 {len(companies)} 家公司,引擎 = {args.engine}")

    engine = get_engine(args.engine)

    def process(name: str) -> dict:
        rec = engine.research(name)
        sc = score_record(rec)
        return {"record": rec, "score": sc}

    results: list[dict] = []
    if args.workers > 1 and len(companies) > 1:
        with ThreadPoolExecutor(max_workers=args.workers) as ex:
            futures = {ex.submit(process, n): n for n in companies}
            for fut in as_completed(futures):
                name = futures[fut]
                try:
                    results.append(fut.result())
                    print(f"  ✓ {name}")
                except Exception as e:  # noqa: BLE001
                    print(f"  ✗ {name}: {e}")
    else:
        for name in companies:
            try:
                results.append(process(name))
                print(f"  ✓ {name}")
            except Exception as e:  # noqa: BLE001
                print(f"  ✗ {name}: {e}")

    if not results:
        sys.exit("没有任何结果。")

    # 排序(overall 降序)
    results.sort(key=lambda r: r["score"]["overall"], reverse=True)
    if args.top:
        results = results[: args.top]

    _write_outputs(results, args.out)


def _write_outputs(results: list[dict], out_dir: str) -> None:
    os.makedirs(out_dir, exist_ok=True)
    reports_dir = os.path.join(out_dir, "reports")
    enriched_dir = os.path.join(out_dir, "enriched")
    os.makedirs(reports_dir, exist_ok=True)
    os.makedirs(enriched_dir, exist_ok=True)

    # 1) 批量 Excel(排名 + 各模块)
    xlsx_path = os.path.join(out_dir, "sales_intelligence.xlsx")
    report.write_excel(results, xlsx_path)

    # 2) 每家公司:完整 Markdown 情报 + 原始 JSON
    for row in results:
        rec, sc = row["record"], row["score"]
        slug = slugify(rec["company_profile"].get("name_en") or "company")
        report.write_markdown(rec, sc, os.path.join(reports_dir, f"{slug}.md"))
        with open(os.path.join(enriched_dir, f"{slug}.json"), "w", encoding="utf-8") as f:
            json.dump(rec, f, ensure_ascii=False, indent=2)

    # 3) 控制台 Top 排名
    print(f"\n🏆 排名(Top {min(len(results), 20)}):")
    print(f"{'#':>2}  {'Tier':<4} {'Overall':>7}  {'Concur':<8} {'公司'}")
    for i, row in enumerate(results[:20], 1):
        sc = row["score"]
        name = row["record"]["company_profile"].get("name_en") or "?"
        print(f"{i:>2}  {sc['tier']:<4} {sc['overall']:>7}  {sc['concur_installed']:<8} {name}")

    print(f"\n✅ 输出:")
    print(f"   • Excel : {xlsx_path}")
    print(f"   • 简报  : {reports_dir}/*.md")
    print(f"   • JSON  : {enriched_dir}/*.json")


def main():
    ap = argparse.ArgumentParser(description="Hitpoint Sales Intelligence Agent")
    ap.add_argument("--company", help="单个公司名")
    ap.add_argument("--input", help="CSV / TSV / XLSX 文件(含公司列)")
    ap.add_argument("--engine", default="manual", help="manual(默认) 或 api")
    ap.add_argument("--out", default="output", help="输出目录(默认 output)")
    ap.add_argument("--top", type=int, default=0, help="只保留前 N 名(0=全部)")
    ap.add_argument("--workers", type=int, default=1, help="并发数(api 引擎建议 3-5)")
    run(ap.parse_args())


if __name__ == "__main__":
    main()
