"""报告渲染:批量 Excel(排名 + 各模块明细)+ 每家公司一份完整 Markdown 情报。

Excel 给销售总监看 Top-N 排序;Markdown 是单客户的 "Account Copilot" 简报。
中英混合:字段名英文,内容中文。
"""

from __future__ import annotations

import os  # noqa: F401  (kept explicit; used below)

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from util import safe_get

# ---------- 样式 ----------
_HEADER_FILL = PatternFill("solid", fgColor="1F3864")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_TIER_FILL = {"A": "C6EFCE", "B": "FFEB9C", "C": "F2F2F2"}


def _join(items, sep="; ") -> str:
    return sep.join(str(x) for x in items if x) if items else ""


def _cities_str(cf: dict) -> str:
    out = []
    for c in cf.get("cities", []) or []:
        hc = f"({c.get('headcount')}人)" if c.get("headcount") else ""
        out.append(f"{c.get('city','')}{hc}")
    return _join(out)


# ======================================================================
# Excel
# ======================================================================

def write_excel(rows: list[dict], path: str) -> None:
    """rows: [{record, score}] 已按 overall 降序。"""
    wb = Workbook()
    _sheet_ranking(wb.active, rows)
    _sheet_profile(wb.create_sheet("公司画像 Profile"), rows)
    _sheet_contacts(wb.create_sheet("采购委员会 Contacts"), rows)
    _sheet_tech(wb.create_sheet("技术栈 TechStack"), rows)
    _sheet_signals(wb.create_sheet("信号 Signals"), rows)
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    wb.save(path)


def _style_header(ws, ncols: int):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"


def _autosize(ws, widths: dict[int, int]):
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w


def _sheet_ranking(ws, rows):
    ws.title = "排名 Top-N"
    headers = ["排名", "公司 Company", "Tier", "Overall", "ICP Fit",
               "Concur", "Cross-Sell", "Buying Intent", "Relationship",
               "中国员工", "AI 总结 (摘要)"]
    ws.append(headers)
    for i, row in enumerate(rows, start=1):
        rec, sc = row["record"], row["score"]
        cp = rec["company_profile"]
        name = cp.get("name_en") or ""
        if cp.get("name_cn"):
            name = f"{name} / {cp['name_cn']}"
        summary = (rec.get("ai_summary_cn") or "").replace("\n", " ")
        ws.append([
            i, name, sc["tier"], sc["overall"], sc["icp_fit"],
            sc["concur_installed"], sc["cross_sell"], sc["buying_intent"],
            sc["relationship"], cp.get("employees_china") or "—",
            summary[:160] + ("…" if len(summary) > 160 else ""),
        ])
        fill = _TIER_FILL.get(sc["tier"])
        if fill:
            ws.cell(row=i + 1, column=3).fill = PatternFill("solid", fgColor=fill)
    _style_header(ws, len(headers))
    _autosize(ws, {1: 6, 2: 32, 3: 6, 4: 9, 5: 8, 6: 9, 7: 11, 8: 13, 9: 13, 10: 12, 11: 70})


def _sheet_profile(ws, rows):
    headers = ["公司", "官网", "中国实体", "中国法人", "母公司", "上市", "股票代码",
               "行业", "Revenue", "全球员工", "中国员工", "中国城市", "实体类型", "共享中心"]
    ws.append(headers)
    for row in rows:
        rec = row["record"]
        cp, cf = rec["company_profile"], rec["china_footprint"]
        ssc = ("有 @" + cf.get("ssc_location", "")) if cf.get("has_shared_service_center") else "—"
        ws.append([
            cp.get("name_en") or "", cp.get("website") or "", cp.get("china_entity_name") or "",
            cp.get("china_legal_rep") or "", cp.get("parent_company") or "",
            "是" if cp.get("is_public") else ("否" if cp.get("is_public") is False else "—"),
            cp.get("ticker") or "—", cp.get("industry") or "", cp.get("revenue") or "",
            cp.get("employees_global") or "—", cp.get("employees_china") or "—",
            _cities_str(cf), _join(cf.get("entity_types", [])), ssc,
        ])
    _style_header(ws, len(headers))
    _autosize(ws, {1: 26, 2: 24, 3: 30, 4: 12, 5: 20, 6: 6, 7: 10, 8: 18, 9: 14, 10: 12, 11: 10, 12: 26, 13: 24, 14: 18})


def _sheet_contacts(ws, rows):
    headers = ["公司", "姓名", "职位 Title", "职能", "Seniority", "Location",
               "在中国", "LinkedIn", "邮箱", "最近换岗", "在岗年限"]
    ws.append(headers)
    for row in rows:
        rec = row["record"]
        company = rec["company_profile"].get("name_en") or ""
        for p in rec.get("buying_committee", []) or []:
            ws.append([
                company, p.get("name") or "", p.get("title") or "", p.get("function") or "",
                p.get("seniority") or "", p.get("location") or "",
                "是" if p.get("in_china") else ("否" if p.get("in_china") is False else "—"),
                p.get("linkedin") or "", p.get("email") or "",
                "是" if p.get("recently_changed_job") else "—", p.get("years_in_role") or "—",
            ])
    _style_header(ws, len(headers))
    _autosize(ws, {1: 22, 2: 14, 3: 30, 4: 14, 5: 12, 6: 14, 7: 8, 8: 36, 9: 24, 10: 10, 11: 10})


def _sheet_tech(ws, rows):
    headers = ["公司", "Concur", "ERP", "费用差旅 T&E", "采购 Procurement",
               "HR", "ITSM", "CRM", "BI/Cloud", "协作", "Concur 详情"]
    ws.append(headers)
    for row in rows:
        rec = row["record"]
        ts = rec["tech_stack"]
        ws.append([
            rec["company_profile"].get("name_en") or "", row["score"]["concur_installed"],
            _join(ts.get("erp", [])), _join(ts.get("expense_travel", [])),
            _join(ts.get("procurement", [])), _join(ts.get("hr", [])),
            _join(ts.get("itsm", [])), _join(ts.get("crm", [])),
            _join(ts.get("bi_cloud", [])), _join(ts.get("collaboration", [])),
            ts.get("concur_details") or "—",
        ])
    _style_header(ws, len(headers))
    _autosize(ws, {1: 22, 2: 9, 3: 22, 4: 20, 5: 20, 6: 16, 7: 14, 8: 16, 9: 20, 10: 20, 11: 30})


def _sheet_signals(ws, rows):
    headers = ["公司", "类型", "日期", "标题/岗位", "摘要/关键词", "触发点", "来源"]
    ws.append(headers)
    for row in rows:
        rec = row["record"]
        company = rec["company_profile"].get("name_en") or ""
        for e in rec.get("business_events", []) or []:
            ws.append([company, f"事件:{e.get('type','')}", e.get("date") or "",
                       e.get("title") or "", (e.get("summary") or "")[:120],
                       "★" if e.get("sales_trigger") else "", e.get("source") or ""])
        for h in rec.get("hiring_signals", []) or []:
            ws.append([company, "招聘", h.get("posted") or "", h.get("title") or "",
                       _join(h.get("keywords", [])), "★" if h.get("transformation_signal") else "",
                       h.get("source") or ""])
    _style_header(ws, len(headers))
    _autosize(ws, {1: 22, 2: 16, 3: 12, 4: 34, 5: 40, 6: 8, 7: 30})


# ======================================================================
# Markdown(单公司完整情报 / Account Copilot 简报)
# ======================================================================

def write_markdown(record: dict, score: dict, path: str) -> None:
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        f.write(render_markdown(record, score))


def render_markdown(rec: dict, sc: dict) -> str:
    cp, cf, ts = rec["company_profile"], rec["china_footprint"], rec["tech_stack"]
    name = cp.get("name_en") or "Unknown"
    cn = f" / {cp['name_cn']}" if cp.get("name_cn") else ""
    L = []
    L.append(f"# {name}{cn} — Company Intelligence\n")
    L.append(f"> Tier **{sc['tier']}** ｜ Overall **{sc['overall']}** ｜ 数据置信度 `{rec.get('data_confidence','?')}` ｜ 引擎 `{rec.get('_engine','?')}`\n")

    # 第十部分 评分(放最前,销售先看结论)
    L.append("## 🎯 Opportunity Score(机会评分)\n")
    L.append("| 维度 | 分数 |")
    L.append("|---|---|")
    L.append(f"| ICP Fit | {sc['icp_fit']} |")
    L.append(f"| Concur Installed | {sc['concur_installed']} |")
    L.append(f"| Cross-Sell | {sc['cross_sell']} |")
    L.append(f"| Buying Intent | {sc['buying_intent']} |")
    L.append(f"| Relationship | {sc['relationship']} |")
    L.append(f"| **Overall** | **{sc['overall']}** |")
    L.append("\n**评分理由:**")
    for dim, reasons in sc["rationale"].items():
        for r in reasons:
            L.append(f"- _{dim}_: {r}")
    L.append("")

    # 第九部分 AI 总结
    if rec.get("ai_summary_cn"):
        L.append("## 🧠 AI Summary(AI 总结)\n")
        L.append(rec["ai_summary_cn"] + "\n")

    # 第一部分 公司画像
    L.append("## 1. Company Profile(公司画像)\n")
    profile_rows = [
        ("官网", cp.get("website")), ("LinkedIn", cp.get("linkedin")),
        ("中国实体", cp.get("china_entity_name")), ("中国法人", cp.get("china_legal_rep")),
        ("成立时间", cp.get("founded_year")), ("母公司", cp.get("parent_company")),
        ("是否上市", "是" if cp.get("is_public") else ("否" if cp.get("is_public") is False else None)),
        ("股票代码", cp.get("ticker")), ("行业", cp.get("industry")),
        ("Revenue", cp.get("revenue")), ("全球员工", cp.get("employees_global")),
        ("中国员工", cp.get("employees_china")), ("总部", cp.get("hq")),
        ("中国办公室", _join(cp.get("china_offices", []))),
    ]
    L.append("| 字段 | 值 |\n|---|---|")
    for k, v in profile_rows:
        if v:
            L.append(f"| {k} | {v} |")
    L.append("")

    # 第二部分 中国业务
    L.append("## 2. China Footprint(中国业务)\n")
    if cf.get("entity_types"):
        L.append(f"**实体类型:** {_join(cf['entity_types'])}")
    if cf.get("has_shared_service_center"):
        L.append(f"**共享服务中心(SSC):** ✅ 位于 {cf.get('ssc_location') or '中国'}")
    if cf.get("cities"):
        L.append("\n| 城市 | 人数 | 职能 |\n|---|---|---|")
        for c in cf["cities"]:
            L.append(f"| {c.get('city','')} | {c.get('headcount') or '—'} | {_join(c.get('functions', []))} |")
    if cf.get("notes"):
        L.append(f"\n{cf['notes']}")
    L.append("")

    # 第三部分 采购委员会
    L.append("## 3. Buying Committee(采购委员会)\n")
    bc = rec.get("buying_committee", []) or []
    if bc:
        L.append("| 姓名 | 职位 | 职能 | Location | 在华 | LinkedIn | 邮箱 |\n|---|---|---|---|---|---|---|")
        for p in bc:
            inc = "✅" if p.get("in_china") else ("—" if p.get("in_china") is False else "?")
            li = f"[link]({p['linkedin']})" if p.get("linkedin") else "—"
            L.append(f"| {p.get('name','')} | {p.get('title','')} | {p.get('function','')} | "
                     f"{p.get('location','')} | {inc} | {li} | {p.get('email') or '—'} |")
    else:
        L.append("_暂无_")
    L.append("")

    # 第四部分 组织结构
    chain = safe_get(rec, "org_chart", "chain", default=[])
    if chain or safe_get(rec, "org_chart", "description"):
        L.append("## 4. Org Chart(组织结构 / 决策链)\n")
        if chain:
            L.append("```\n" + "\n  ↓\n".join(chain) + "\n```")
        if safe_get(rec, "org_chart", "description"):
            L.append(rec["org_chart"]["description"])
        L.append("")

    # 第五部分 技术栈
    L.append("## 5. Technology Stack(技术栈)\n")
    tech_cats = [("ERP", "erp"), ("费用/差旅 T&E", "expense_travel"), ("采购 Procurement", "procurement"),
                 ("HR", "hr"), ("ITSM", "itsm"), ("CRM", "crm"), ("BI/Cloud", "bi_cloud"),
                 ("协作", "collaboration")]
    L.append("| 类别 | 系统 |\n|---|---|")
    for label, key in tech_cats:
        if ts.get(key):
            L.append(f"| {label} | {_join(ts[key])} |")
    L.append(f"| **Concur** | **{sc['concur_installed']}** {('— ' + ts['concur_details']) if ts.get('concur_details') else ''} |")
    if ts.get("details"):
        L.append("\n<details><summary>技术栈细节(来源/置信度)</summary>\n")
        L.append("\n| 类别 | 厂商 | 产品 | Cloud | 起始 | Partner | 置信度 | 来源 |\n|---|---|---|---|---|---|---|---|")
        for d in ts["details"]:
            L.append(f"| {d.get('category','')} | {d.get('vendor','')} | {d.get('product','')} | "
                     f"{'是' if d.get('cloud') else '—'} | {d.get('since') or '—'} | "
                     f"{d.get('partner') or '—'} | {d.get('confidence') or '—'} | {d.get('source') or '—'} |")
        L.append("\n</details>")
    L.append("")

    # 第六部分 商业事件
    events = rec.get("business_events", []) or []
    if events:
        L.append("## 6. Business Events(商业事件 / Sales Triggers)\n")
        for e in events:
            star = " 🔥" if e.get("sales_trigger") else ""
            L.append(f"- **{e.get('date','')}** [{e.get('type','')}] {e.get('title','')}{star}  \n  {e.get('summary','')}"
                     + (f"  \n  来源: {e['source']}" if e.get("source") else ""))
        L.append("")

    # 第七部分 招聘信号
    hires = rec.get("hiring_signals", []) or []
    if hires:
        L.append("## 7. Hiring Signals(招聘信号)\n")
        for h in hires:
            star = " 🔥" if h.get("transformation_signal") else ""
            L.append(f"- **{h.get('title','')}**{star} — {h.get('location','')} {h.get('posted','')}  \n  "
                     f"关键词: {_join(h.get('keywords', []))}"
                     + (f"  \n  来源: {h['source']}" if h.get("source") else ""))
        L.append("")

    # Next Best Action(MVP 第七模块)
    L.append("## 8. Next Best Action(下一步建议)\n")
    for action in next_best_actions(rec, sc):
        L.append(f"- {action}")
    L.append("")

    # 来源
    if rec.get("sources"):
        L.append("## 📚 Sources(信息来源)\n")
        for s in rec["sources"]:
            L.append(f"- {s}")
        L.append("")

    return "\n".join(L)


def next_best_actions(rec: dict, sc: dict) -> list[str]:
    """根据评分自动生成下一步动作建议。"""
    acts = []
    bc = rec.get("buying_committee", []) or []
    cf = rec["china_footprint"]

    if sc["cross_sell"] >= 80 and sc["concur_installed"] != "Yes":
        acts.append("**主切入点**:该公司跑大型 ERP 但(很可能)未上 Concur —— 以「SAP + Concur 一体化费控」为主线切入。")
    elif sc["concur_installed"] == "Yes":
        acts.append("已用 Concur —— 切入点转向托管运维 / 优化 / 增购模块(发票、采购、审计合规)。")

    # 找切入人:优先 SSC 负责人,其次 Finance Director / CFO
    ssc_people = [p for p in bc if "shared" in (p.get("function", "") + p.get("title", "")).lower()
                  or "共享" in (p.get("title", ""))]
    fin_people = [p for p in bc if (p.get("function", "").lower() == "finance") or "cfo" in p.get("title", "").lower()
                  or "finance director" in p.get("title", "").lower()]
    target = (ssc_people or fin_people or bc)[:1]
    if target:
        t = target[0]
        who = "共享服务负责人" if ssc_people else "财务决策人"
        acts.append(f"**首要联系人**({who}):{t.get('name','')} — {t.get('title','')}"
                    + (f"({t['linkedin']})" if t.get("linkedin") else "") + "。")
    else:
        acts.append("**待补**:尚未锁定采购委员会关键人 —— 优先补齐 Shared Service / Finance Director / IT(SAP)三类联系人。")

    triggers = [e for e in rec.get("business_events", []) or [] if e.get("sales_trigger")]
    if triggers:
        acts.append(f"**Why now**:{triggers[0].get('title','')} —— 以此事件作为开场由头。")
    transform_hires = [h for h in rec.get("hiring_signals", []) or [] if h.get("transformation_signal")]
    if transform_hires:
        acts.append(f"**信号**:正在招「{transform_hires[0].get('title','')}」,说明转型在推进,时机好。")

    if cf.get("has_shared_service_center"):
        acts.append(f"**话术**:围绕 {cf.get('ssc_location') or ''} 共享中心的费用合规、发票自动化、跨实体报销标准化展开。")

    acts.append("**下一步**:生成定制开场邮件 + LinkedIn 连接语 + Demo 议程(可由 Agent 自动起草)。")
    return acts
